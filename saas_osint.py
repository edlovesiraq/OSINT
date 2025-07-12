
import socket
import re

def is_valid_domain(domain):
    if len(domain) > 253:
        return False
    if any(len(label) > 63 for label in domain.split(".")):
        return False
    pattern = re.compile(r"^(?!-)[A-Za-z0-9.-]{1,253}(?<!-)$")
    return bool(pattern.match(domain))

def resolve_domains_to_ips(domains):
    """Resolve subdomains to IP addresses (skip malformed/oversized)."""
    ips = set()
    for domain in domains:
        domain = domain.strip().lower()
        if not is_valid_domain(domain):
            print(f"[!] Skipping invalid domain: {domain}")
            continue
        try:
            ip = socket.gethostbyname(domain)
            ips.add(ip)
        except Exception as e:
            print(f"[!] Could not resolve {domain}: {e}")
    return list(ips)


import os
import subprocess
import requests
import json
import time
from datetime import datetime
import webbrowser
from pathlib import Path
import argparse
import signal
import urllib.parse
import sys

# Configuration
TOOLS = {
    'amass': 'amass',
    'theHarvester': 'theHarvester',
    'git': 'git'
}

# Colors for console output
class Colors:
    HEADER = '\033[95m'
    OKBLUE = '\033[94m'
    OKGREEN = '\033[92m'
    WARNING = '\033[93m'
    FAIL = '\033[91m'
    ENDC = '\033[0m'
    BOLD = '\033[1m'
    UNDERLINE = '\033[4m'

def print_banner():
    banner = f"""
{Colors.HEADER}{Colors.BOLD}
  ____  _____ ____  _____   _____ _____ _   _ _______ 
 / ___||_   _/ ___||_   _| |_   _| ____| /| |_   _|
 \___ / | | \___ / | |     | | |  _| |  \| | | |  
  ___) | | |  ___) | | |     | | | |___| |/ | | |  
 |____/  |_| |____/  |_|     |_| |_____|_| \_| |_|  
                                                     
{Colors.ENDC}
{Colors.OKBLUE}SaaS OSINT Automation Tool v3{Colors.ENDC}
{Colors.OKGREEN}Passive reconnaissance for SaaS targets{Colors.ENDC}
"""
    print(banner)

def check_tools():
    missing_tools = []
    for tool, cmd in TOOLS.items():
        try:
            # Special check for theHarvester on Kali
            if tool == 'theHarvester':
                result = subprocess.run(['which', 'theHarvester'], stdout=subprocess.PIPE, stderr=subprocess.PIPE)
                if result.returncode != 0:
                    result = subprocess.run(['which', 'theharvester'], stdout=subprocess.PIPE, stderr=subprocess.PIPE)
                    if result.returncode == 0:
                        TOOLS['theHarvester'] = 'theharvester'
                    else:
                        missing_tools.append(tool)
            else:
                subprocess.run([cmd, '--version'], stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=True)
        except (subprocess.CalledProcessError, FileNotFoundError):
            missing_tools.append(tool)
    
    if missing_tools:
        print(f"{Colors.FAIL}Error: The following tools are not installed or not in PATH:{Colors.ENDC}")
        for tool in missing_tools:
            print(f"- {tool}")
        print("\nPlease install them before running this script.")
        print("For Kali Linux, try:")
        print("sudo apt update && sudo apt install amass theharvester git")
        exit(1)

def create_workspace(target):
    workspace = Path(f"saas-osint-{target}")
    directories = [
        "recon",
        "emails",
        "leaks",
        "screenshots",
        "certs",
        "reports",
        "shodan"
    ]
    
    print(f"{Colors.OKBLUE}[*]{Colors.ENDC} Creating workspace for {target}")
    
    try:
        workspace.mkdir(exist_ok=True)
        for directory in directories:
            (workspace / directory).mkdir(exist_ok=True)
        print(f"{Colors.OKGREEN}[+]{Colors.ENDC} Workspace created at {workspace}")
    except Exception as e:
        print(f"{Colors.FAIL}[-]{Colors.ENDC} Error creating workspace: {e}")
        exit(1)
    
    return workspace

def run_amass(target, workspace, timeout=10):
    print(f"{Colors.OKBLUE}[*]{Colors.ENDC} Running Amass for passive subdomain discovery (timeout: {timeout} minutes)")
    output_file = workspace / "recon" / "amass_domains.txt"
    
    try:
        # Run Amass with timeout
        process = subprocess.Popen([
            TOOLS['amass'],
            'enum',
            '-passive',
            '-d', target,
            '-o', str(output_file),
            '-timeout', str(timeout)
        ])
        
        try:
            process.wait(timeout=timeout * 60 + 30)  # Add 30 seconds buffer
        except subprocess.TimeoutExpired:
            print(f"{Colors.WARNING}[!]{Colors.ENDC} Amass timed out after {timeout} minutes")
            process.kill()
        
        domains = []
        if output_file.exists():
            with open(output_file, 'r') as f:
                domains = [line.strip() for line in f if line.strip()]
        
        print(f"{Colors.OKGREEN}[+]{Colors.ENDC} Found {len(domains)} subdomains")
        return domains
    except Exception as e:
        print(f"{Colors.FAIL}[-]{Colors.ENDC} Amass failed: {e}")
        return []

def query_crtsh(target, workspace):
    print(f"{Colors.OKBLUE}[*]{Colors.ENDC} Querying crt.sh for certificate transparency logs")
    url = f"https://crt.sh/?q=%.{target}&output=json"
    output_file = workspace / "certs" / "crtsh_results.json"
    
    try:
        response = requests.get(url, timeout=30)
        response.raise_for_status()
        data = response.json()
        
        with open(output_file, 'w') as f:
            json.dump(data, f, indent=2)
        
        unique_domains = set()
        for entry in data:
            name_value = entry.get('name_value', '')
            if name_value:
                for domain in name_value.split('\n'):
                    if domain.strip() and target in domain:
                        unique_domains.add(domain.strip().lower())
        
        print(f"{Colors.OKGREEN}[+]{Colors.ENDC} Found {len(unique_domains)} unique domains from certificates")
        return list(unique_domains)
    except Exception as e:
        print(f"{Colors.FAIL}[-]{Colors.ENDC} crt.sh query failed: {e}")
        return []


def run_theharvester(target, workspace):
    print(f"{Colors.OKBLUE}[*]{Colors.ENDC} Running theHarvester for email and employee data")
    output_file = workspace / "recon" / "harvest_report.xml"
    email_output = workspace / "emails" / "harvested_emails.txt"
    
    try:
        subprocess.run([
            TOOLS['theHarvester'],
            '-d', target,
            '-b', 'baidu,bing,duckduckgo,linkedin,netcraft,otx,threatcrowd,yahoo',
            '-f', str(output_file.with_suffix(''))
        ], check=True)
        
        # Parse XML for email addresses
        import xml.etree.ElementTree as ET
        emails = set()
        try:
            tree = ET.parse(output_file)
            root = tree.getroot()
            for elem in root.iter('email'):
                email = elem.text.strip()
                if email:
                    emails.add(email)
        except Exception as e:
            print(f"{Colors.WARNING}[!]{Colors.ENDC} Could not parse emails from XML: {e}")
        
        if emails:
            with open(email_output, 'w') as f:
                f.write("")
f.write("\n".join(sorted(emails)))print(f"{Colors.OKGREEN}[+]{Colors.ENDC} Extracted {len(emails)} emails from theHarvester")
        else:
            print(f"{Colors.WARNING}[!]{Colors.ENDC} No emails extracted from theHarvester")
        
        return True
    except subprocess.CalledProcessError as e:
        print(f"{Colors.FAIL}[-]{Colors.ENDC} theHarvester failed: {e}")
        return False

def open_hunter_io(target):
    print(f"{Colors.OKBLUE}[*]{Colors.ENDC} Opening Hunter.io in browser for email pattern discovery")
    url = f"https://hunter.io/search/{target}"
    webbrowser.open(url)
    time.sleep(2)  # Rate limiting

def open_linkedin_search(target):
    print(f"{Colors.OKBLUE}[*]{Colors.ENDC} Opening LinkedIn search in browser")
    query = f"site:linkedin.com/in \"{target}\" OR \"{target.replace('-', ' ')}\""
    url = f"https://www.google.com/search?q={urllib.parse.quote(query)}"
    webbrowser.open(url)
    time.sleep(2)  # Rate limiting

def run_shodan_search(target, workspace):
    print(f"{Colors.OKBLUE}[*]{Colors.ENDC} Running Shodan search")
    
    # First try with API key if available
    shodan_api_key = os.getenv('SHODAN_API_KEY', '')
    if shodan_api_key:
        try:
            from shodan import Shodan
            api = Shodan(shodan_api_key)
            
            # Search for the target
            results = api.search(f'ssl:"{target}" OR hostname:"{target}" OR org:"{target}"')
            
            # Save raw results
            with open(workspace / "shodan" / "shodan_results.json", 'w') as f:
                json.dump(results, f, indent=2)
            
            # Collect unique IPs
            ips = set()
            for result in results.get('matches', []):
                ips.add(result['ip_str'])
            
            # Save IP list
            with open(workspace / "shodan" / "shodan_ips.txt", 'w') as f:
                f.write("\n".join(ips))print(f"{Colors.OKGREEN}[+]{Colors.ENDC} Found {len(ips)} Shodan results")
            return list(ips)
        except Exception as e:
            print(f"{Colors.WARNING}[!]{Colors.ENDC} Shodan API search failed: {e}")
    
    # Fall back to manual search
    print(f"{Colors.WARNING}[!]{Colors.ENDC} Opening Shodan in browser for manual search")
    queries = [
        f'ssl:"{target}"',
        f'hostname:"{target}"',
        f'org:"{target}"'
    ]
    for query in queries:
        url = f"https://www.shodan.io/search?query={urllib.parse.quote(query)}"
        webbrowser.open(url)
        time.sleep(2)
    return []

def open_censys_search(target):
    print(f"{Colors.OKBLUE}[*]{Colors.ENDC} Opening Censys search in browser")
    queries = [
        target,
        f'names: {target}',
        f'443.https.tls.certificate.parsed.names: {target}'
    ]
    for query in queries:
        url = f"https://search.censys.io/search?q={urllib.parse.quote(query)}"
        webbrowser.open(url)
        time.sleep(2)  # Rate limiting

def check_haveibeenpwned(emails_file, workspace):
    if not emails_file.exists():
        print(f"{Colors.WARNING}[!]{Colors.ENDC} No email file found to check HaveIBeenPwned")
        return
    
    print(f"{Colors.OKBLUE}[*]{Colors.ENDC} Opening HaveIBeenPwned in browser")
    url = "https://haveibeenpwned.com/"
    webbrowser.open(url)
    time.sleep(1)

def run_github_dorking(target, workspace):
    print(f"{Colors.OKBLUE}[*]{Colors.ENDC} Preparing GitHub dorking")
    output_file = workspace / "leaks" / "github_secrets.txt"
    
    # Simple GitHub search without API
    queries = [
        f'{target} filename:.env',
        f'{target} filename:config.yml',
        f'{target} password',
        f'{target} api_key',
        f'{target} secret',
        f'{target} credentials'
    ]
    
    try:
        with open(output_file, 'w') as f:
            for query in queries:
                url = f"https://github.com/search?q={urllib.parse.quote(query)}&type=code"
                f.write(f"Search: {query}\n")
                f.write(f"URL: {url}\n\n")
                webbrowser.open(url)
                time.sleep(5)  # Longer delay for GitHub searches
        
        print(f"{Colors.OKGREEN}[+]{Colors.ENDC} GitHub dorking completed (manual review required)")
        return True
    except Exception as e:
        print(f"{Colors.FAIL}[-]{Colors.ENDC} GitHub dorking failed: {e}")
        return False

def generate_google_dorks(target):
    dorks = [
        f'site:{target} ext:pdf',
        f'site:{target} intitle:index.of',
        f'"password" filetype:log site:{target}',
        f'inurl:admin site:{target}',
        f'site:{target} (intext:"username" OR intext:"password")',
        f'site:{target} ext:sql|xls|xlsx|doc|docx|ppt|pptx',
        f'site:{target} filetype:env OR filetype:ini OR filetype:cfg',
        f'site:{target} "api key" OR "secret key" OR "private key"'
    ]
    
    print(f"{Colors.OKBLUE}[*]{Colors.ENDC} Generated Google dorks for manual searching:")
    for i, dork in enumerate(dorks):
        url = f"https://www.google.com/search?q={urllib.parse.quote(dork)}"
        print(f"  {Colors.BOLD}{dork}{Colors.ENDC}")
        webbrowser.open(url)
        # Progressive delay to avoid rate limiting
        time.sleep(5 + i)

def generate_report(target, workspace, domains, cert_domains, shodan_ips):
    report_file = workspace / "reports" / f"{target}_osint_report.md"
    
    with open(report_file, 'w') as f:
        f.write(f"# OSINT Report for {target}\n\n")
        f.write(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")
        
        f.write("## Domain Enumeration\n\n")
        f.write("### Subdomains from Amass\n")
        f.write("```\n")
        f.write("\n".join(domains) if domains else "No subdomains found")
        f.write("\n```\n\n")
        
        f.write("### Subdomains from Certificate Transparency\n")
        f.write("```\n")
        f.write("\n".join(cert_domains) if cert_domains else "No certificate domains found")
        f.write("\n```\n\n")
        
        if shodan_ips:
            f.write("## Shodan Results\n\n")
            f.write(f"Found {len(shodan_ips)} IP addresses with services matching the target\n")
            f.write("```\n")
            f.write("\n".join(shodan_ips))f.write("\n```\n\n")
            f.write("Detailed reports available in the shodan/ directory\n")
        
        f.write("## Recommended Next Steps\n")
        f.write("- Review theHarvester output in recon/harvest_report.xml\n")
        f.write("- Check Hunter.io for email patterns\n")
        f.write("- Review GitHub dorking results in leaks/github_secrets.txt\n")
        f.write("- Follow up on Shodan/Censys findings manually\n")
        f.write("- Check Google dorks results manually\n")
    
    print(f"{Colors.OKGREEN}[+]{Colors.ENDC} Report generated at {report_file}")

def main():
    parser = argparse.ArgumentParser(description='Automated SaaS OSINT Recon Tool v3')
    parser.add_argument('target', help='Target domain (e.g., example-saas.com)')
    parser.add_argument('--timeout', type=int, default=10, 
                      help='Timeout in minutes for Amass scan (default: 10)')
    args = parser.parse_args()
    
    print_banner()
    check_tools()
    
    target = args.target.lower()
    workspace = create_workspace(target)
    
    # Phase 2: Domain Enumeration
    domains = run_amass(target, workspace, args.timeout)
    cert_domains = query_crtsh(target, workspace)
    
    # Phase 3: Email & Employee Data
    run_theharvester(target, workspace)
    open_hunter_io(target)
    open_linkedin_search(target)
    
    # Phase 4: Infrastructure Fingerprinting
    resolved_ips = resolve_domains_to_ips(domains + cert_domains)
    shodan_ips = run_shodan_host_lookup(resolved_ips, workspace)
    open_censys_search(target)
    
    # Phase 5: Data Leak & Breach Checks
    check_haveibeenpwned(workspace / "emails" / "harvested_emails.txt", workspace)
    run_github_dorking(target, workspace)
    
    # Phase 6: Deep Web & Google Dorking
    generate_google_dorks(target)
    
    # Generate report
    generate_report(target, workspace, domains, cert_domains, shodan_ips)
    
    print(f"\n{Colors.OKGREEN}{Colors.BOLD}[+] OSINT collection complete!{Colors.ENDC}")
    print(f"Review the findings in {workspace}")

if __name__ == "__main__":
    main()


import requests
from docx import Document
from docx.shared import Inches

def write_word_report(target, workspace, domains, resolved_ips, shodan_results, harvested_emails, hibp_hits):
    doc = Document()
    doc.add_heading(f"OSINT Recon Report for {target}", 0)

    doc.add_heading("Subdomains", level=1)
    for d in domains:
        doc.add_paragraph(d)

    doc.add_heading("Resolved IPs", level=1)
    for ip in resolved_ips:
        doc.add_paragraph(ip)

    doc.add_heading("Shodan Results", level=1)
    for ip, data in shodan_results.items():
        doc.add_paragraph(f"{ip}: {data.get('hostnames', [])} - {data.get('org', 'N/A')}")

    doc.add_heading("Harvested Emails", level=1)
    for email in harvested_emails:
        doc.add_paragraph(email)

    doc.add_heading("HIBP Hits", level=1)
    for email, breaches in hibp_hits.items():
        doc.add_paragraph(f"{email} - Breaches: {', '.join(breaches) if breaches else 'None'}")

    output_path = workspace / "recon" / "OSINT_Report.docx"
    doc.save(output_path)
    print(f"[+] Word report saved to {output_path}")

def check_hibp_api(emails, hibp_key):
    headers = {"hibp-api-key": hibp_key, "user-agent": "osint-checker"}
    url_template = "https://haveibeenpwned.com/api/v3/breachedaccount/{}"
    results = {}
    for email in emails:
        try:
            res = requests.get(url_template.format(email), headers=headers)
            if res.status_code == 200:
                results[email] = [entry["Name"] for entry in res.json()]
            elif res.status_code == 404:
                results[email] = []
            else:
                print(f"[!] HIBP API error for {email}: {res.status_code}")
                results[email] = None
        except Exception as e:
            print(f"[!] HIBP check failed for {email}: {e}")
            results[email] = None
    return results
